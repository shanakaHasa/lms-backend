"""Document parsing: bytes in, page-aware text out.

Page numbers are carried all the way through to the citation, so a teacher can
open the source PDF at the page an answer came from. That single requirement is
why parsing returns pages rather than one flat string.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass

from app.core.errors import ValidationFailed
from app.core.logging import get_logger

log = get_logger(__name__)

SUPPORTED_MIME = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "text/plain": "txt",
    "text/markdown": "txt",
}


@dataclass(frozen=True)
class Page:
    number: int
    text: str


@dataclass(frozen=True)
class ParsedDocument:
    pages: list[Page]
    page_count: int
    detected_type: str

    @property
    def is_empty(self) -> bool:
        return not any(p.text.strip() for p in self.pages)


# Course materials are heavily templated. These are the natural section
# boundaries in a syllabus, an assignment brief or an assessment policy, and
# cutting on them is what stops "Late penalties" being glued onto the tail of
# "Assessment" -- see the chunker for why that matters.
_HEADING = re.compile(
    r"^\s*(?:"
    r"[A-Z][A-Z \t/&'-]{3,60}"  # ALL CAPS heading
    # "1.2 Assessment", "1.2.3 Late Penalties", "2. Overview". The trailing dot
    # is optional: real documents write "1.2 Title" far more often than "1.2.".
    r"|\d+(?:\.\d+){0,3}\.?[ \t]+[A-Z][\w \t/&'-]{2,60}"
    # "Week 3", "Week 3 - Linked Lists", "Module 2: Recursion". Deliberately
    # NOT `Week\s+\d+.*`: that matches any sentence beginning "Week 1 ...",
    # which turns ordinary syllabus prose into headings and shreds the chunks.
    # Excluding sentence punctuation from the tail is what draws the line.
    # The en and em dashes are deliberate: real course documents use both.
    r"|(?:Week|Module|Topic)\s+\d+[ \t]*(?:[-–—:][ \t]*)?[\w \t/&'-]{0,50}"  # noqa: RUF001
    r"|(?:Course|Unit|Subject|Overview|Description|Aims?|Objectives?|"
    r"Learning\s+Outcomes?|Prerequisites?|Assessment|Assessments?|Grading|"
    r"Marking|Schedule|Timetable|Readings?|Resources?|Textbooks?|"
    r"Late\s+Penalties|Extensions?|Special\s+Consideration|"
    r"Academic\s+Integrity|Plagiarism|Attendance|Contact|Staff|"
    r"Submission|Deliverables?|Rubric|Weighting|Due\s+Dates?)\s*:?"
    r")\s*$",
    re.MULTILINE,
)


def detect_type(mime_type: str, filename: str) -> str:
    if mime_type in SUPPORTED_MIME:
        return SUPPORTED_MIME[mime_type]
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix in {"pdf", "docx", "txt", "md"}:
        return "txt" if suffix == "md" else suffix
    raise ValidationFailed(
        f"unsupported document type: {mime_type or filename}",
        detail={"supported": sorted(SUPPORTED_MIME)},
    )


def parse(content: bytes, mime_type: str, filename: str) -> ParsedDocument:
    kind = detect_type(mime_type, filename)
    if kind == "pdf":
        return _parse_pdf(content)
    if kind == "docx":
        return _parse_docx(content)
    return _parse_text(content)


def _parse_pdf(content: bytes) -> ParsedDocument:
    import pymupdf

    pages: list[Page] = []
    with pymupdf.open(stream=content, filetype="pdf") as doc:
        for i, page in enumerate(doc, start=1):
            # "text" preserves reading order well enough for course documents.
            # A scanned handout needs OCR, which is a separate concern.
            pages.append(Page(number=i, text=_normalise(page.get_text("text"))))
    parsed = ParsedDocument(pages=pages, page_count=len(pages), detected_type="pdf")
    if parsed.is_empty:
        # A scan looks like a valid PDF and produces nothing. Failing loudly
        # here beats indexing an empty document and silently never retrieving it.
        raise ValidationFailed("PDF contained no extractable text (likely a scan; OCR is required)")
    return parsed


def _parse_docx(content: bytes) -> ParsedDocument:
    import docx

    document = docx.Document(io.BytesIO(content))
    parts: list[str] = [p.text for p in document.paragraphs]
    # Assessment weightings and schedules are almost always tables; dropping
    # them would lose exactly the facts people ask about.
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    # DOCX has no page concept without rendering; treat the file as one page.
    return ParsedDocument(
        pages=[Page(number=1, text=_normalise("\n".join(parts)))],
        page_count=1,
        detected_type="docx",
    )


def _parse_text(content: bytes) -> ParsedDocument:
    text = _normalise(content.decode("utf-8", errors="replace"))
    return ParsedDocument(pages=[Page(number=1, text=text)], page_count=1, detected_type="txt")


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\xa0", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def find_sections(text: str) -> list[tuple[int, str]]:
    """Return (offset, heading) pairs so chunks can be labelled with a section."""
    return [(m.start(), m.group().strip()) for m in _HEADING.finditer(text)]
